"""B2 测试：文档加载与切块（模块 1）。

覆盖三件事：
1. 四种格式（PDF / DOCX / TXT / MD）都能加载，且 metadata 统一；
2. PDF 必须带页码（引用展示依赖它）；TXT/MD 无页码；
3. 切块参数（500/50）生效，块内去重生效，chunk_id 稳定可复现。

测试数据在 ``tmp_path`` 中现场生成，不依赖仓库里的样例文档，
这样即使样例被改动也不会让测试变红。
"""

from __future__ import annotations

from pathlib import Path

import pytest


# ---------------------------------------------------------------------------
# 测试数据
# ---------------------------------------------------------------------------
LONG_TEXT = "\n\n".join(
    [
        "第一条 年假规定：员工入职满一年后享有年假，工作满一年不满十年者每年五天。",
        "第二条 病假规定：员工因病需要休息的，凭二级以上医院开具的病假证明申请病假。",
        "第三条 事假规定：事假为无薪假，全年累计不超过十五天，需提前三个工作日申请。",
        "第四条 报销标准：一线城市住宿标准为每晚六百元，其他城市为每晚四百元。",
        "第五条 保密义务：员工在职期间接触到的技术资料与客户名单均属公司商业秘密。",
        "第六条 离职交接：员工主动离职需提前三十日以书面形式通知公司。",
        "第七条 加班调休：法定节假日加班按三倍工资支付，或安排等额调休。",
        "第八条 培训支持：公司为员工提供每年三千元的技能培训补贴。",
    ]
)


@pytest.fixture()
def txt_file(tmp_path: Path) -> Path:
    """UTF-8 中文 TXT。"""
    path = tmp_path / "制度.txt"
    path.write_text(LONG_TEXT, encoding="utf-8")
    return path


@pytest.fixture()
def gbk_file(tmp_path: Path) -> Path:
    """GBK 编码 TXT（中文 Windows 上极常见的"乱码源"）。"""
    path = tmp_path / "旧制度.txt"
    path.write_bytes(LONG_TEXT.encode("gb18030"))
    return path


@pytest.fixture()
def md_file(tmp_path: Path) -> Path:
    """Markdown 文件。"""
    path = tmp_path / "FAQ.md"
    path.write_text(f"# 运维 FAQ\n\n## 备份策略\n\n{LONG_TEXT}\n", encoding="utf-8")
    return path


@pytest.fixture()
def docx_file(tmp_path: Path) -> Path:
    """DOCX 文件（用 python-docx 现场生成）。"""
    from docx import Document

    path = tmp_path / "产品需求.docx"
    document = Document()
    document.add_heading("云笔记 V2.0 产品需求", level=0)
    document.add_heading("功能范围", level=1)
    for line in LONG_TEXT.split("\n\n"):
        document.add_paragraph(line)
    document.save(str(path))
    return path


@pytest.fixture()
def pdf_file(tmp_path: Path) -> Path:
    """PDF 文件（用 reportlab 生成两页中文内容，页码断言依赖它）。"""
    reportlab = pytest.importorskip("reportlab", reason="生成 PDF 测试数据需要 reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    # 找一个可用的中文 TrueType 字体；找不到就跳过该用例（CI 可能没有中文字体）
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    ]
    font_path = next((item for item in candidates if item.exists()), None)
    if font_path is None:
        pytest.skip("系统缺少中文字体，跳过 PDF 用例")

    font_name = "TestCN"
    if font_path.suffix.lower() == ".ttc":
        pdfmetrics.registerFont(TTFont(font_name, str(font_path), subfontIndex=0))
    else:
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))

    path = tmp_path / "员工手册.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    for page_index, chunk in enumerate([LONG_TEXT[: len(LONG_TEXT) // 2], LONG_TEXT[len(LONG_TEXT) // 2 :]]):
        pdf.setFont(font_name, 12)
        text_object = pdf.beginText(40, 800)
        for line in chunk.split("\n\n"):
            text_object.textLine(line)
        pdf.drawText(text_object)
        pdf.showPage()
    pdf.save()
    return path


# ---------------------------------------------------------------------------
# 加载
# ---------------------------------------------------------------------------
def test_load_txt(txt_file: Path) -> None:
    """TXT 加载：metadata 统一，无页码。"""
    from rag.loaders import load_file

    outcome = load_file(txt_file)
    assert outcome.ok is True
    assert outcome.doc_id and len(outcome.doc_id) == 16
    assert outcome.char_count > 100
    assert outcome.page_count is None
    assert outcome.degraded is False

    document = outcome.documents[0]
    assert document.metadata["file_name"] == "制度.txt"
    assert document.metadata["ext"] == ".txt"
    assert document.metadata["doc_id"] == outcome.doc_id
    assert "page" not in document.metadata
    assert "年假" in document.page_content


def test_load_gbk_txt_is_not_mojibake(gbk_file: Path) -> None:
    """GBK 编码的 TXT 必须能正确解码，并标记为编码降级。"""
    from rag.loaders import load_file

    outcome = load_file(gbk_file)
    assert outcome.ok is True
    assert outcome.degraded is True          # 走了 GB18030 回退，必须如实标记
    assert "年假" in outcome.documents[0].page_content
    assert "锟斤拷" not in outcome.documents[0].page_content


def test_load_markdown(md_file: Path) -> None:
    """Markdown 加载。"""
    from rag.loaders import load_file

    outcome = load_file(md_file)
    assert outcome.ok is True
    assert outcome.documents[0].metadata["ext"] == ".md"
    assert "运维 FAQ" in outcome.documents[0].page_content


def test_load_docx(docx_file: Path) -> None:
    """DOCX 加载（Docx2txtLoader 路径）。"""
    from rag.loaders import load_file

    outcome = load_file(docx_file)
    assert outcome.ok is True
    assert outcome.documents[0].metadata["file_name"] == "产品需求.docx"
    assert "云笔记" in outcome.documents[0].page_content
    assert outcome.page_count is None        # DOCX 不提供页码


def test_load_pdf_has_page_numbers(pdf_file: Path) -> None:
    """PDF 加载：每页一个 Document，页码从 1 开始（引用"第几页"依赖此断言）。"""
    from rag.loaders import load_file

    outcome = load_file(pdf_file)
    assert outcome.ok is True
    assert outcome.page_count == 2
    assert len(outcome.documents) == 2

    pages = sorted(doc.metadata["page"] for doc in outcome.documents)
    assert pages == [1, 2]
    for document in outcome.documents:
        assert document.metadata["file_name"] == "员工手册.pdf"
        assert document.metadata["doc_id"] == outcome.doc_id


def test_load_rejects_unsupported_format(tmp_path: Path) -> None:
    """不支持的格式必须给出明确错误，而不是抛异常。"""
    from rag.loaders import load_file

    path = tmp_path / "data.xlsx"
    path.write_bytes(b"fake")
    outcome = load_file(path)
    assert outcome.ok is False
    assert outcome.error_type == "unsupported_format"
    assert "不支持" in (outcome.error or "")


def test_load_missing_file(tmp_path: Path) -> None:
    """文件不存在 → 结构化错误。"""
    from rag.loaders import load_file

    outcome = load_file(tmp_path / "nope.txt")
    assert outcome.ok is False
    assert outcome.error_type == "file_not_found"


def test_load_empty_file(tmp_path: Path) -> None:
    """空文件 → empty_document。"""
    from rag.loaders import load_file

    path = tmp_path / "空.txt"
    path.write_text("   \n  \n", encoding="utf-8")
    outcome = load_file(path)
    assert outcome.ok is False
    assert outcome.error_type == "empty_document"


def test_doc_id_is_content_based(txt_file: Path, tmp_path: Path) -> None:
    """doc_id 由内容决定：改名不影响，改内容才变化。"""
    from rag.loaders import compute_doc_id, load_file

    copied = tmp_path / "改名后.txt"
    copied.write_bytes(txt_file.read_bytes())
    assert compute_doc_id(txt_file) == compute_doc_id(copied)
    assert load_file(txt_file).doc_id == load_file(copied).doc_id

    changed = tmp_path / "改内容.txt"
    changed.write_text(LONG_TEXT + "新增第九条。", encoding="utf-8")
    assert compute_doc_id(changed) != compute_doc_id(txt_file)


def test_load_directory(tmp_path: Path) -> None:
    """目录批量加载：忽略不支持格式，返回每个文件的结果。"""
    from rag.loaders import load_directory

    (tmp_path / "a.txt").write_text(LONG_TEXT, encoding="utf-8")
    (tmp_path / "b.md").write_text(LONG_TEXT, encoding="utf-8")
    (tmp_path / "c.xlsx").write_bytes(b"ignored")
    outcomes = load_directory(tmp_path)
    assert len(outcomes) == 2                       # xlsx 被忽略
    assert all(item.ok for item in outcomes)


# ---------------------------------------------------------------------------
# 切块
# ---------------------------------------------------------------------------
def test_split_respects_chunk_size(txt_file: Path) -> None:
    """chunk_size=500 生效：除极短尾块外都不超过 500 字符。"""
    from config.settings import get_settings
    from rag.loaders import load_file
    from rag.splitter import split_documents

    settings = get_settings()
    chunks, _ = split_documents(load_file(txt_file).documents, settings)

    assert settings.chunk_size == 500
    assert settings.chunk_overlap == 50
    assert len(chunks) >= 1
    assert max(chunk.char_len for chunk in chunks) <= settings.chunk_size + 5


def test_split_metadata_and_chunk_id(tmp_path: Path) -> None:
    """切块必须保留元数据，并生成稳定可复现的 chunk_id。"""
    from rag.loaders import load_file
    from rag.splitter import split_documents

    path = tmp_path / "长文.txt"
    path.write_text(LONG_TEXT * 6, encoding="utf-8")   # 足够长，切出多块

    outcome = load_file(path)
    chunks, duplicates = split_documents(outcome.documents)

    assert len(chunks) >= 3
    assert duplicates >= 0
    for index, chunk in enumerate(chunks):
        assert chunk.file_name == "长文.txt"
        assert chunk.doc_id == outcome.doc_id
        assert chunk.chunk_id.startswith(f"{outcome.doc_id}:0:")
        assert chunk.page is None
        assert chunk.char_len == len(chunk.text)
        assert index == chunk.chunk_index

    # 同一输入两次切分结果完全一致（chunk_id 稳定，才能做幂等写入）
    chunks_again, _ = split_documents(outcome.documents)
    assert [c.chunk_id for c in chunks] == [c.chunk_id for c in chunks_again]


def test_split_deduplicates_identical_blocks() -> None:
    """内容完全相同的块只保留一份（避免同一段话挤占 top_k）。

    构造：把同一份 Document 输入两次（模拟 PDF 里重复的页面、或同一文件被重复加载），
    这样切块阶段一定会产出内容相同的块，去重逻辑必须把它们压成一份。
    """
    from langchain_core.documents import Document

    from rag.splitter import split_documents

    paragraph = (
        "公司的年假政策为：入职满一年享有五天年假，满十年享有十天，满二十年享有十五天；"
        "当年未休完的部分最多结转三天至次年三月三十一日前使用完毕。"
    )
    document = Document(
        page_content=paragraph,
        metadata={"file_name": "重复.txt", "doc_id": "docdup", "ext": ".txt", "page": 1},
    )

    # 第一次：正常切块
    single, _ = split_documents([document])
    assert len(single) == 1

    # 第二次：同一内容出现两份 → 必须去重成一份
    deduped, duplicates = split_documents([document, document])
    assert len(deduped) == 1, "重复内容必须只保留一份"
    assert duplicates >= 1, "重复块必须被计入 duplicates 统计"

    # 去重后不存在内容相同的块
    normalized = [chunk.text.strip() for chunk in deduped]
    assert len(normalized) == len(set(normalized))


def test_split_preserves_page_for_pdf(pdf_file: Path) -> None:
    """PDF 切块后每个块都要带页码，否则无法给出精确引用。"""
    from rag.loaders import load_file
    from rag.splitter import split_documents

    chunks, _ = split_documents(load_file(pdf_file).documents)
    assert chunks
    assert all(chunk.page in (1, 2) for chunk in chunks)
    assert {chunk.page for chunk in chunks} == {1, 2}


def test_split_text_helper() -> None:
    """直接切分一段文本（便于测试与粘贴文本场景）。"""
    from rag.splitter import split_text

    chunks = split_text(LONG_TEXT, file_name="t.txt", doc_id="abc")
    assert chunks
    assert chunks[0].doc_id == "abc"
    assert chunks[0].chunk_id.startswith("abc:0:")


def test_splitter_summary() -> None:
    """切分参数摘要可读取（供 /health 与调试展示）。"""
    from rag.splitter import splitter_summary

    summary = splitter_summary()
    assert summary["chunk_size"] == 500
    assert summary["chunk_overlap"] == 50
    assert summary["separators"][0] == "\n\n"
    assert "。" in summary["separators"]


# ---------------------------------------------------------------------------
# 文件名清洗（安全）
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "raw,forbidden",
    [
        ("../../etc/passwd.txt", "/"),
        (r"..\..\windows\system32\cmd.txt", "\\"),
        ("a<b>c:d.txt", "<"),
        ("  空格名.txt", " "),
    ],
)
def test_sanitize_filename_blocks_path_traversal(raw: str, forbidden: str) -> None:
    """文件名清洗必须挡住路径穿越与特殊字符。"""
    from rag.models import sanitize_filename

    cleaned = sanitize_filename(raw)
    assert forbidden not in cleaned
    assert ".." not in cleaned
    assert cleaned.endswith(".txt")


def test_sanitize_filename_keeps_chinese() -> None:
    """中文文件名要保留，不能因为清洗变成乱码。"""
    from rag.models import sanitize_filename

    assert sanitize_filename("员工手册（2025版）.pdf") == "员工手册（2025版）.pdf"


def test_sanitize_filename_fallback() -> None:
    """全是非法字符时要有兜底名。"""
    from rag.models import sanitize_filename

    assert sanitize_filename("!!!") == "document"
    assert sanitize_filename("") == "document"
